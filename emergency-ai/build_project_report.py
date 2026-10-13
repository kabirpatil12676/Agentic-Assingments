from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR / "output"
OUT_DIR.mkdir(exist_ok=True)
OUT_FILE = OUT_DIR / "Emergency_Response_CrewAI_Project_Report.docx"
TEMP_DIR = Path(r"C:\Users\Lenovo\AppData\Local\Temp")

EVIDENCE = [
    ("Planner Agent output", "codex-clipboard-fad1e61e-1f34-4575-be23-38eced07a0d6.png"),
    ("Patient Analysis Agent output", "codex-clipboard-637ddd89-3e57-4989-bd03-2e59755dcef4.png"),
    ("Resource Allocation Agent output", "codex-clipboard-2b72e078-63d8-485a-84db-d52330a55bed.png"),
    ("Ambulance Agent output", "codex-clipboard-5274be2a-2c24-4e4b-9b57-c63fe84cbb38.png"),
    ("Surgery Scheduling Agent output", "codex-clipboard-cb4291cd-4fee-4bf2-aff6-6d4f18cf99d0.png"),
    ("Monitoring Agent output", "codex-clipboard-654743f3-59c5-4637-8d39-4681e48de356.png"),
    ("Learning Agent output", "codex-clipboard-9a20c6f9-d1eb-4bb2-bb26-4dbcd5de2e82.png"),
]


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), color)
    tc_pr.append(shade)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths=None):
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[cell_index])
            if row_index == 0:
                set_cell_shading(cell, "1F4E78")
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead:
        paragraph.add_run(bold_lead).bold = True
    paragraph.add_run(text)
    return paragraph


def add_bullets(doc, entries):
    for entry in entries:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(entry)


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_evidence_page(doc, title, image_name, figure_number):
    doc.add_page_break()
    add_heading(doc, "Appendix A Run Evidence", 1)
    add_heading(doc, title, 2)
    path = TEMP_DIR / image_name
    if path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(6.65))
        add_caption(doc, f"Figure {figure_number}. Captured final answer from the {title}.")
    else:
        add_body(doc, "Evidence image was unavailable when this report was generated.")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer.paragraphs[0]
    set_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Emergency Response CrewAI Project Report")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Seven Agent Simulation Output Review and Validation").italic = True
    subtitle.paragraph_format.space_after = Pt(16)

    add_body(doc, "This report documents the Emergency Response CrewAI simulation, evaluates the captured outputs from its seven agents, and identifies the controls needed before the workflow can be considered reliable for an educational demonstration. The run completed across all seven stages, and the supplied mock resource facts were generally retained. However, five outputs did not fully follow their own safety or formatting instructions. The overall result is therefore conditionally acceptable for an assignment demonstration, but not fully compliant with the current prompts.")

    summary = doc.add_table(rows=1, cols=2)
    summary.rows[0].cells[0].text = "Report item"
    summary.rows[0].cells[1].text = "Value"
    for label, value in [
        ("Project", "Emergency Response CrewAI simulation"),
        ("Workflow", "Seven CrewAI agents executed sequentially"),
        ("Test case", "Road-accident simulation with severe leg bleeding and severe pain"),
        ("Evidence reviewed", "Thirteen screenshots covering prompts and final agent answers"),
        ("Overall output status", "Partially compliant; 2 pass, 5 need prompt or validation fixes"),
        ("Scope", "Educational decision-support simulation only; not real emergency dispatch or clinical care"),
        ("Report date", "12 October 2026"),
    ]:
        cells = summary.add_row().cells
        cells[0].text = label
        cells[1].text = value
    format_table(summary, [2.0, 4.7])

    add_heading(doc, "1 Project Purpose Objectives and Scope")
    add_body(doc, "The project demonstrates how CrewAI can coordinate specialist roles around one simulated emergency case. It uses a planner, triage analyst, resource allocator, ambulance selector, surgery scheduler, monitoring analyst, and learning agent. The Assignment 3 report was used as the structural reference: it organizes a seven-agent workflow around purpose, architecture, evidence, quality findings, limitations, and recommendations.")
    add_heading(doc, "Objectives", 2)
    add_bullets(doc, [
        "Break a simulated emergency into sequenced coordination tasks.",
        "Summarize patient urgency without turning the agent into a treatment provider.",
        "Ground ambulance, bed, clinician, and theatre decisions in supplied mock data.",
        "Record monitoring trends and quality-improvement lessons from anonymized mock cases.",
        "Produce a repeatable report while preserving the simulation-only boundary.",
    ])
    add_heading(doc, "Scope boundaries", 2)
    add_body(doc, "The system contains no connection to emergency services, hospital systems, patient records, live ambulance tracking, theatre booking, or clinical devices. Its outputs must be treated as generated coursework evidence only. A live deployment would require clinical governance, verified data interfaces, audit logging, human authorization, cybersecurity controls, and formal safety validation.")

    add_heading(doc, "2 System Design and Technology")
    add_body(doc, "The project follows the same high-level pattern as Assignment 3: deterministic mock data is assembled in Python, concise summaries are supplied to CrewAI tasks, agents execute in sequence, and the final results are saved as a local report. Gemini is configured through CrewAI with a low temperature and delegation disabled so that each role remains narrow and inspectable.")
    design = doc.add_table(rows=1, cols=3)
    for cell, text in zip(design.rows[0].cells, ["Layer", "Implementation", "Responsibility"]):
        cell.text = text
    for row in [
        ("Mock data", "Python modules in tools", "Ambulance status, clinicians, beds, operating theatres, vital trends, and historical lessons"),
        ("Orchestration", "CrewAI sequential process", "Passes prior task outputs as context to later specialist agents"),
        ("Language model", "Gemini through CrewAI LLM", "Produces constrained natural-language planning and analysis"),
        ("Configuration", "python-dotenv", "Loads the API key and optional model override"),
        ("Output", "Python text report", "Stores the completed simulated response locally"),
    ]:
        cells = design.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    format_table(design, [1.15, 1.85, 3.7])

    add_heading(doc, "3 Seven Agent Architecture")
    agents = doc.add_table(rows=1, cols=3)
    for cell, text in zip(agents.rows[0].cells, ["Agent", "Primary responsibility", "Evidence based finding"]):
        cell.text = text
    for row in [
        ("Planner", "Creates five ordered coordination steps", "Produced five steps, but included prohibited treatment language"),
        ("Patient Analysis", "Classifies urgency and prepares resources", "Critical classification was coherent, but output included diagnoses and treatment directions"),
        ("Resource Allocation", "Selects mock trauma resources", "Used available trauma bay, one resuscitation bed, named specialists, and OT-2"),
        ("Ambulance", "Selects an eligible mock ambulance", "Correctly chose AMB-12, the nearest available advanced-life-support unit"),
        ("Surgery Scheduling", "Prepares earliest mock theatre option", "Correctly selected OT-2 at 14:30"),
        ("Monitoring", "Summarizes supplied vital signs", "Described the supplied trend accurately, but included a care instruction"),
        ("Learning", "Suggests process improvements", "Produced three relevant lessons, but one recommendation reads as a clinical directive"),
    ]:
        cells = agents.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    format_table(agents, [1.35, 2.05, 3.3])

    add_heading(doc, "4 End to End Workflow")
    add_bullets(doc, [
        "Load the emergency scenario and concise mock availability and monitoring summaries.",
        "Run the Planner and Patient Analysis tasks to establish simulated urgency and coordination needs.",
        "Run Resource Allocation and Ambulance tasks against the supplied availability facts.",
        "Use the triage and resource context to prepare the earliest mock operating-theatre option.",
        "Interpret the supplied vital-sign snapshots and derive process lessons from anonymized prior cases.",
        "Save the seven outputs to a local emergency-response simulation report.",
    ])

    add_heading(doc, "5 Test Method and Captured Evidence")
    add_body(doc, "The review compared the final answers visible in the supplied screenshots against each task contract and the source facts shown in the prompts. This is an output-contract review, not a clinical review. A task is marked Pass only when it meets the stated structure and safety requirements; Partial means the principal data decision was correct but one or more task constraints were broken.")
    validation = doc.add_table(rows=1, cols=4)
    for cell, text in zip(validation.rows[0].cells, ["Agent", "Status", "Correct elements", "Issue found"]):
        cell.text = text
    for row in [
        ("Planner", "Partial", "Exactly five ordered coordination steps", "Contains treatment language such as stabilization and life-saving interventions despite the prohibition"),
        ("Patient Analysis", "Partial", "Critical acuity and immediate transport priority fit the simulation", "Names diagnoses and treatment items despite the instruction not to diagnose or prescribe"),
        ("Resource Allocation", "Pass", "Exactly four bullets; uses supplied trauma bay, bed, doctors, OT-2, and 14:30", "No material contract breach observed"),
        ("Ambulance", "Partial", "Correctly selects available AMB-12 with 7-minute ETA and advanced capability", "Does not begin with the required Selected ambulance phrase"),
        ("Surgery Scheduling", "Pass", "Uses the earliest supplied available slot: OT-2 at 14:30", "No material contract breach observed"),
        ("Monitoring", "Partial", "Exactly three bullets; numerical trend matches supplied snapshots", "Prescribes immediate advanced-life-support intervention, which the task prohibited"),
        ("Learning", "Partial", "Exactly three relevant process improvements with no patient identifiers", "Mandatory surgical evaluation language can be read as clinical prescribing"),
    ]:
        cells = validation.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    format_table(validation, [1.1, 0.75, 2.25, 2.6])

    add_heading(doc, "6 Detailed Output Review")
    add_heading(doc, "6.1 Decisions grounded in the supplied mock data", 2)
    add_body(doc, "The three operational selections are consistent with the data visible in the run. AMB-12 is the nearest available unit and is equipped for advanced life support; AMB-03 is closer but marked busy. The resource allocator correctly identified the available trauma bay, resuscitation-bed capacity, named clinicians, and OT-2. The surgery scheduler then chose the earliest theatre time, 14:30, rather than the later 15:15 option.")
    add_heading(doc, "6.2 Output-contract violations", 2)
    add_body(doc, "The principal weakness is prompt compliance, not the sequencing of agents. The Planner, Patient Analysis, and Monitoring outputs crossed the stated boundary against treatment instructions. The Learning output is useful at a process level but should avoid phrases that require a clinical decision. In addition, the Ambulance output reached the correct selection but missed the required opening format. These findings mean that the run cannot be described as fully correct under the current task definitions.")
    add_heading(doc, "6.3 Consistency across agents", 2)
    add_body(doc, "The outputs are internally consistent: the triage stage describes a high-risk simulated case; the resource and ambulance stages prepare an appropriate simulated response; the theatre recommendation matches the availability table; and the monitoring output refers to AMB-12 selected earlier in the workflow. The remaining concern is that generated language can add clinical content that was neither required nor supplied by the mock data.")

    add_heading(doc, "7 Recommendations")
    add_bullets(doc, [
        "Replace clinical terms in task outputs with neutral coordination wording, for example notify, prepare, confirm, hand over, and escalate to an authorized clinician.",
        "Add a programmatic guard that flags prohibited phrases such as diagnose, stabilize, transfusion, intervention, treatment, analgesia, and mandatory surgical evaluation before a report is saved.",
        "Validate the Ambulance task structure with a prefix check requiring Selected ambulance as the first text in the response.",
        "Use a structured schema for each task output so fields such as ambulance ID, ETA, theatre, and availability can be checked deterministically.",
        "Keep all availability data explicitly mock and prevent integrations with real dispatch, hospital, or patient systems in the assignment version.",
    ])

    add_heading(doc, "8 Conclusion")
    add_body(doc, "The project successfully demonstrates a connected seven-agent CrewAI workflow and the key resource decisions in the captured run are correct relative to the supplied mock data. Its task-compliance result is partial rather than full: Resource Allocation and Surgery Scheduling pass, while Planner, Patient Analysis, Ambulance, Monitoring, and Learning each need at least one prompt or validation improvement. With deterministic output checks and stricter non-clinical wording, the project will be stronger and more defensible as an assignment-ready simulation.")

    for figure_number, (title, image_name) in enumerate(EVIDENCE, start=1):
        add_evidence_page(doc, title, image_name, figure_number)

    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build()
